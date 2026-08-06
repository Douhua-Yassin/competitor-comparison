from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

Document = pytest.importorskip("docx").Document

from app.reporting import analysis_engine
from app.reporting.dashboard_db import (
    list_settings_products,
    save_note,
    update_listing_scope,
    upsert_daily_metric,
    upsert_listings,
    upsert_stores,
)
from app.reporting.report_archive import list_report_runs
from app.reporting.report_generator import generate_report
from app.reporting.targets import import_targets


def _write_target(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "目标"
    sheet.append([
        "周期类型", "周期", "产品线", "店铺SID", "ASIN", "MSKU",
        "销售额目标", "销量目标", "利润目标", "利润率目标",
        "TACOS目标", "广告花费目标", "FBA可售库存目标", "备注",
    ])
    sheet.append(["日", "2026-08-04", "足球门", "", "", "", 500, 8, "", "", "10%", 50, 100, "日报目标"])
    workbook.save(path)


def test_rule_based_word_report_contains_targets_charts_and_archive(tmp_path: Path, monkeypatch):
    db_path = tmp_path / "reporting.db"
    upsert_stores([{"sid": 10, "name": "美国店", "country": "美国"}], db_path)
    upsert_listings(
        [
            {
                "sid": 10,
                "asin": "B000TEST01",
                "msku": "GOAL-01",
                "product_name": "足球门 6x4",
                "country": "美国",
                "status": 1,
            }
        ],
        db_path,
    )
    product = list_settings_products(db_path)["products"][0]
    update_listing_scope(product["id"], "key", "足球门", db_path)
    for metric_date, sales, spend, stock in (
        ("2026-08-02", 200, 25, 120),
        ("2026-08-03", 260, 28, 115),
        ("2026-08-04", 300, 30, 110),
    ):
        upsert_daily_metric(metric_date, product["id"], "sales_amount", sales, "currency", "orders", db_path=db_path)
        upsert_daily_metric(metric_date, product["id"], "units", sales / 50, "count", "orders", db_path=db_path)
        upsert_daily_metric(metric_date, product["id"], "ad_spend", spend, "currency", "sp_product_report", db_path=db_path)
        upsert_daily_metric(metric_date, product["id"], "fba_available", stock, "count", "fba_inventory", db_path=db_path)
    save_note(
        "足球门",
        "day",
        "2026-08-04:2026-08-04",
        "今天降低了广告预算。",
        db_path,
    )
    target_path = tmp_path / "目标表.xlsx"
    _write_target(target_path)
    import_targets(target_path, db_path)
    monkeypatch.setattr(
        analysis_engine,
        "deepseek_settings",
        lambda root: {"api_key": "", "base_url": "", "model": "", "timeout": 90},
    )

    path = generate_report(
        "足球门",
        "day",
        output_dir=tmp_path / "reports",
        today=date(2026, 8, 4),
        db_path=db_path,
    )

    assert path.exists()
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "足球门经营日报" in text
    assert "今天降低了广告预算" in text
    assert "目标完成情况" in text
    assert "销售额" in table_text
    assert "500.00" in table_text
    assert len(document.tables) >= 3
    assert len(document.inline_shapes) >= 1

    history = list_report_runs("足球门", db_path=db_path)
    assert len(history) == 1
    assert history[0]["artifact_sha256"]
    assert history[0]["analysis_source"] == "rules"
