from __future__ import annotations

import hashlib
import math
import os
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .dashboard_db import DB_PATH, ROOT, WRITE_LOCK, connection

DEFAULT_TARGET_FILE = ROOT / "目标表.xlsx"

METRIC_COLUMNS: dict[str, tuple[str, str, str, str]] = {
    "销售额目标": ("sales_amount", "currency", "higher", "销售额"),
    "销量目标": ("units", "count", "higher", "销量"),
    "利润目标": ("profit", "currency", "higher", "利润"),
    "利润率目标": ("profit_margin", "ratio", "higher", "利润率"),
    "TACOS目标": ("tacos", "ratio", "lower", "TACOS"),
    "广告花费目标": ("ad_spend", "currency", "budget", "广告花费"),
    "FBA可售库存目标": ("fba_available", "count", "neutral", "FBA可售库存"),
}
PERIOD_ALIASES = {
    "日": "day",
    "日报": "day",
    "day": "day",
    "周": "week",
    "周报": "week",
    "week": "week",
    "月": "month",
    "月报": "month",
    "month": "month",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def target_file_path() -> Path:
    configured = os.getenv("REPORT_TARGETS_FILE", "").strip()
    if not configured:
        return DEFAULT_TARGET_FILE
    path = Path(configured)
    return path if path.is_absolute() else ROOT / path


def init_targets_db(db_path: Optional[Path] = None) -> None:
    path = Path(db_path or DB_PATH)
    with WRITE_LOCK:
        with connection(path) as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS target_imports (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  source_path TEXT NOT NULL,
                  source_hash TEXT NOT NULL,
                  imported_at TEXT NOT NULL,
                  status TEXT NOT NULL,
                  row_count INTEGER NOT NULL DEFAULT 0,
                  message TEXT
                );

                CREATE TABLE IF NOT EXISTS report_targets (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  import_id INTEGER NOT NULL,
                  version INTEGER NOT NULL,
                  active INTEGER NOT NULL DEFAULT 1,
                  period_type TEXT NOT NULL
                    CHECK(period_type IN ('day','week','month')),
                  period_key TEXT NOT NULL,
                  period_start TEXT NOT NULL,
                  period_end TEXT NOT NULL,
                  scope_type TEXT NOT NULL
                    CHECK(scope_type IN ('product_line','listing')),
                  product_line TEXT NOT NULL,
                  sid INTEGER,
                  asin TEXT,
                  msku TEXT,
                  metric_code TEXT NOT NULL,
                  target_value REAL NOT NULL,
                  unit TEXT NOT NULL,
                  direction TEXT NOT NULL,
                  note TEXT,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(import_id) REFERENCES target_imports(id)
                );

                CREATE INDEX IF NOT EXISTS idx_report_targets_lookup
                  ON report_targets(product_line, period_type, period_start, period_end, active);
                CREATE INDEX IF NOT EXISTS idx_report_targets_scope
                  ON report_targets(scope_type, sid, asin, msku, metric_code, active);
                """
            )


def write_target_template(path: Optional[Path] = None) -> Path:
    output = Path(path or target_file_path())
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    guide_rows = [
        ["字段", "说明"],
        ["周期类型", "填写：日、周、月"],
        ["周期", "日：2026-08-06；周：2026-W32；月：2026-08"],
        ["产品线", "必填，例如：足球门"],
        ["店铺SID/ASIN/MSKU", "全部留空表示产品线目标；填写任一项表示单品目标"],
        ["百分比", "可以填写 10% 或 0.1"],
        ["导入规则", "同一范围、周期和指标再次导入时，旧版本停用但保留历史"],
    ]
    for row in guide_rows:
        guide.append(row)

    sheet = workbook.create_sheet("目标")
    headers = [
        "周期类型",
        "周期",
        "产品线",
        "店铺SID",
        "ASIN",
        "MSKU",
        *METRIC_COLUMNS.keys(),
        "备注",
    ]
    sheet.append(headers)
    sheet.append(["月", "2026-08", "足球门", "", "", "", 50000, 600, 8000, "16%", "12%", 7000, 300, "月度目标示例"])
    sheet.append(["周", "2026-W32", "足球门", "", "", "", 12000, 150, "", "", "13%", 1800, "", "周目标示例"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5EFF")
    sheet.freeze_panes = "A2"
    widths = {
        "A": 12, "B": 15, "C": 20, "D": 12, "E": 16, "F": 18,
        "G": 14, "H": 12, "I": 12, "J": 14, "K": 12, "L": 14, "M": 18, "N": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    try:
        workbook.save(output)
    finally:
        workbook.close()
    return output


def import_targets(
    source_path: Optional[Path] = None,
    db_path: Optional[Path] = None,
) -> dict[str, Any]:
    path = Path(source_path or target_file_path())
    database = Path(db_path or DB_PATH)
    init_targets_db(database)
    if not path.exists():
        raise ValueError(f"未找到目标表：{path}")

    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    with connection(database) as db:
        duplicate = db.execute(
            "SELECT * FROM target_imports WHERE source_hash=? AND status='success' ORDER BY id DESC LIMIT 1",
            (source_hash,),
        ).fetchone()
        if duplicate:
            return {
                "imported": False,
                "duplicate": True,
                "row_count": int(duplicate["row_count"]),
                "source_path": str(path),
                "source_hash": source_hash,
                "imported_at": duplicate["imported_at"],
            }

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["目标"] if "目标" in workbook.sheetnames else next(
            (workbook[name] for name in workbook.sheetnames if name != "填写说明"),
            workbook.active,
        )
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("目标表没有内容")
    header = [_normalize_header(value) for value in rows[0]]
    positions = {name: index for index, name in enumerate(header) if name}
    for required in ("周期类型", "周期", "产品线"):
        if required not in positions:
            raise ValueError(f"目标表缺少必填列：{required}")

    parsed: list[dict[str, Any]] = []
    errors: list[str] = []
    for excel_row, values in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue
        try:
            period_type, period_key, period_start, period_end = _parse_period(
                _value(values, positions, "周期类型"),
                _value(values, positions, "周期"),
            )
            product_line = _text(_value(values, positions, "产品线"))
            if not product_line:
                raise ValueError("产品线不能为空")
            sid = _int_or_none(_value(values, positions, "店铺SID"))
            asin = _text(_value(values, positions, "ASIN"))
            msku = _text(_value(values, positions, "MSKU"))
            scope_type = "listing" if any((sid is not None, asin, msku)) else "product_line"
            note = _text(_value(values, positions, "备注"))
            metric_count = 0
            for column, (code, unit, direction, _label) in METRIC_COLUMNS.items():
                if column not in positions:
                    continue
                raw = _value(values, positions, column)
                if raw in (None, ""):
                    continue
                metric_count += 1
                parsed.append(
                    {
                        "period_type": period_type,
                        "period_key": period_key,
                        "period_start": period_start,
                        "period_end": period_end,
                        "scope_type": scope_type,
                        "product_line": product_line,
                        "sid": sid,
                        "asin": asin,
                        "msku": msku,
                        "metric_code": code,
                        "target_value": _number(raw, ratio=unit == "ratio"),
                        "unit": unit,
                        "direction": direction,
                        "note": note,
                    }
                )
            if metric_count == 0:
                raise ValueError("没有填写任何目标值")
        except ValueError as exc:
            errors.append(f"第{excel_row}行：{exc}")
    if errors:
        preview = "；".join(errors[:10])
        if len(errors) > 10:
            preview += f"；另有{len(errors) - 10}行错误"
        raise ValueError(preview)
    if not parsed:
        raise ValueError("目标表没有可导入的目标")

    imported_at = utc_now()
    with WRITE_LOCK:
        with connection(database) as db:
            cursor = db.execute(
                """
                INSERT INTO target_imports (
                  source_path, source_hash, imported_at, status, row_count, message
                ) VALUES (?, ?, ?, 'running', 0, NULL)
                """,
                (str(path), source_hash, imported_at),
            )
            import_id = int(cursor.lastrowid)
            try:
                for item in parsed:
                    key = (
                        item["period_type"],
                        item["period_key"],
                        item["scope_type"],
                        item["product_line"],
                        item["sid"],
                        item["asin"],
                        item["msku"],
                        item["metric_code"],
                    )
                    version = int(
                        db.execute(
                            """
                            SELECT COALESCE(MAX(version), 0) + 1
                            FROM report_targets
                            WHERE period_type=? AND period_key=? AND scope_type=?
                              AND product_line=?
                              AND COALESCE(sid, -1)=COALESCE(?, -1)
                              AND COALESCE(asin, '')=COALESCE(?, '')
                              AND COALESCE(msku, '')=COALESCE(?, '')
                              AND metric_code=?
                            """,
                            key,
                        ).fetchone()[0]
                    )
                    db.execute(
                        """
                        UPDATE report_targets SET active=0
                        WHERE period_type=? AND period_key=? AND scope_type=?
                          AND product_line=?
                          AND COALESCE(sid, -1)=COALESCE(?, -1)
                          AND COALESCE(asin, '')=COALESCE(?, '')
                          AND COALESCE(msku, '')=COALESCE(?, '')
                          AND metric_code=? AND active=1
                        """,
                        key,
                    )
                    db.execute(
                        """
                        INSERT INTO report_targets (
                          import_id, version, active, period_type, period_key,
                          period_start, period_end, scope_type, product_line,
                          sid, asin, msku, metric_code, target_value, unit,
                          direction, note, created_at
                        ) VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            import_id,
                            version,
                            item["period_type"],
                            item["period_key"],
                            item["period_start"],
                            item["period_end"],
                            item["scope_type"],
                            item["product_line"],
                            item["sid"],
                            item["asin"],
                            item["msku"],
                            item["metric_code"],
                            item["target_value"],
                            item["unit"],
                            item["direction"],
                            item["note"],
                            imported_at,
                        ),
                    )
                db.execute(
                    "UPDATE target_imports SET status='success', row_count=?, message=? WHERE id=?",
                    (len(parsed), "导入成功", import_id),
                )
            except Exception as exc:
                db.execute(
                    "UPDATE target_imports SET status='failed', message=? WHERE id=?",
                    (f"{type(exc).__name__}: {exc}", import_id),
                )
                raise

    return {
        "imported": True,
        "duplicate": False,
        "row_count": len(parsed),
        "source_path": str(path),
        "source_hash": source_hash,
        "imported_at": imported_at,
    }


def target_status(db_path: Optional[Path] = None) -> dict[str, Any]:
    database = Path(db_path or DB_PATH)
    init_targets_db(database)
    path = target_file_path()
    with connection(database) as db:
        latest = db.execute(
            "SELECT * FROM target_imports ORDER BY id DESC LIMIT 1"
        ).fetchone()
        active_count = int(
            db.execute("SELECT COUNT(*) FROM report_targets WHERE active=1").fetchone()[0]
        )
    return {
        "source_path": str(path),
        "source_exists": path.exists(),
        "active_target_count": active_count,
        "last_import": dict(latest) if latest else None,
    }


def targets_for_period(
    product_line: str,
    period_type: str,
    start: str,
    end: str,
    summary: dict[str, Any],
    db_path: Optional[Path] = None,
    *,
    products: Optional[list[dict[str, Any]]] = None,
    metric_rows: Optional[list[dict[str, Any]]] = None,
) -> list[dict[str, Any]]:
    """Backward-compatible entry point for listing-aware target evaluation."""
    from .target_evaluation import targets_for_period as evaluate_targets

    return evaluate_targets(
        product_line,
        period_type,
        start,
        end,
        summary,
        db_path,
        products=products,
        metric_rows=metric_rows,
    )



def _parse_period(period_type_value: Any, period_value: Any) -> tuple[str, str, str, str]:
    alias = _text(period_type_value).lower()
    period_type = PERIOD_ALIASES.get(alias)
    if period_type is None:
        raise ValueError("周期类型必须是日、周或月")

    if isinstance(period_value, datetime):
        day_value = period_value.date()
    elif isinstance(period_value, date):
        day_value = period_value
    else:
        day_value = None
    text = _text(period_value)

    if period_type == "day":
        current = day_value or _parse_date(text)
        return period_type, current.isoformat(), current.isoformat(), current.isoformat()
    if period_type == "week":
        if day_value:
            current = day_value
            year, week, _ = current.isocalendar()
        else:
            match = re.fullmatch(r"(\d{4})-?W(\d{1,2})", text, flags=re.IGNORECASE)
            if not match:
                current = _parse_date(text)
                year, week, _ = current.isocalendar()
            else:
                year, week = int(match.group(1)), int(match.group(2))
        start = date.fromisocalendar(year, week, 1)
        end = start + timedelta(days=6)
        return period_type, f"{year}-W{week:02d}", start.isoformat(), end.isoformat()

    if day_value:
        year, month = day_value.year, day_value.month
    else:
        match = re.fullmatch(r"(\d{4})[-/](\d{1,2})", text)
        if not match:
            current = _parse_date(text)
            year, month = current.year, current.month
        else:
            year, month = int(match.group(1)), int(match.group(2))
    start = date(year, month, 1)
    next_month = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    end = next_month - timedelta(days=1)
    return period_type, f"{year}-{month:02d}", start.isoformat(), end.isoformat()


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"无法识别周期：{value}") from exc


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value))


def _value(values: tuple[Any, ...], positions: dict[str, int], name: str) -> Any:
    index = positions.get(name)
    return values[index] if index is not None and index < len(values) else None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _int_or_none(value: Any) -> Optional[int]:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"店铺SID不是整数：{value}") from exc


def _number(value: Any, *, ratio: bool = False) -> float:
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if text.endswith("%"):
            try:
                return float(text[:-1]) / 100
            except ValueError as exc:
                raise ValueError(f"无法识别数值：{value}") from exc
        try:
            number = float(text)
        except ValueError as exc:
            raise ValueError(f"无法识别数值：{value}") from exc
    else:
        try:
            number = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"无法识别数值：{value}") from exc
    if not math.isfinite(number):
        raise ValueError(f"目标值必须是有限数字：{value}")
    if ratio and abs(number) > 1:
        return number / 100
    return number
