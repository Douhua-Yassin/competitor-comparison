from __future__ import annotations

import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .analysis_engine import generate_analysis
from .charts import create_report_charts
from .dashboard_db import DATA_DIR, DB_PATH
from .dashboard_service import report_context
from .docx_style import configure_document_styles
from .report_archive import archive_report

ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = DATA_DIR / "reports"
REPORT_LABEL = {"day": "日报", "week": "周报", "month": "月报"}


def generate_report(
    product_line: str,
    report_type: str,
    *,
    output_dir: Optional[Path] = None,
    today: Optional[date] = None,
    db_path: Optional[Path] = None,
    analysis_post: Any = None,
) -> Path:
    if report_type not in REPORT_LABEL:
        raise ValueError("报告类型必须是 day、week 或 month")
    current = today or date.today()
    database = Path(db_path or DB_PATH)
    module, selected = report_context(product_line, report_type, database, current)
    fallback = _fallback_analysis(module, selected)
    context = {
        "report_type": report_type,
        "product_line": module["product_line"],
        "report_period": {"start": selected["start"], "end": selected["end"]},
        "selected_window": selected,
        "targets": selected.get("targets") or [],
        "all_window_notes": {
            item["code"]: (item.get("note") or {}).get("content", "")
            for item in module.get("windows") or []
        },
        "products": module["products"],
    }
    analysis_result = generate_analysis(context, fallback, ROOT, post=analysis_post)
    analysis = analysis_result.analysis

    target_dir = Path(output_dir or REPORT_DIR / current.isoformat())
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_line = re.sub(r'[\\/:*?"<>|]+', "_", product_line).strip() or "产品线"
    stamp = datetime.now().strftime("%H%M%S-%f")
    path = target_dir / f"{safe_line}-{selected['start']}-{selected['end']}-{REPORT_LABEL[report_type]}-{stamp}.docx"

    document = Document()
    configure_document_styles(document)
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    document.core_properties.title = f"{product_line}经营{REPORT_LABEL[report_type]}"
    document.core_properties.subject = f"{selected['start']} 至 {selected['end']}"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"{product_line}经营{REPORT_LABEL[report_type]}")
    title_run.bold = True
    title_run.font.size = Pt(18)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"报告周期：{selected['start']} 至 {selected['end']}")
    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_run = source.add_run(
        "分析来源：DeepSeek结构化分析" if analysis_result.source == "deepseek" else "分析来源：程序规则版"
    )
    source_run.italic = True
    if analysis_result.warning:
        paragraph = document.add_paragraph()
        paragraph.add_run("分析提示：" + analysis_result.warning).italic = True

    document.add_heading("一、领导摘要", level=1)
    _add_bullets(document, analysis.get("executive_summary"))

    document.add_heading("二、核心数据与目标", level=1)
    _add_summary_table(document, selected)
    _add_target_table(document, selected.get("targets") or [])
    if selected.get("source_note"):
        paragraph = document.add_paragraph()
        paragraph.add_run("数据口径说明：" + selected["source_note"]).italic = True
    for warning in selected.get("warnings") or []:
        document.add_paragraph("数据提示：" + warning)

    document.add_heading("三、趋势图", level=1)
    with tempfile.TemporaryDirectory(prefix="report-charts-", dir=target_dir) as chart_dir:
        charts = create_report_charts(selected.get("series") or {}, Path(chart_dir), safe_line)
        if not charts:
            document.add_paragraph("当前周期没有足够的逐日数据生成趋势图。")
        for chart_path, caption in charts:
            document.add_picture(str(chart_path), width=Cm(16.5))
            caption_paragraph = document.add_paragraph(caption)
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("四、趋势与变化", level=1)
    _add_analysis_section(document, "表现判断", analysis.get("performance"))
    _add_analysis_section(document, "变化原因", analysis.get("causes"))

    document.add_heading("五、已记录的操作与判断", level=1)
    note_snapshot: dict[str, str] = {}
    any_note = False
    for window in module.get("windows") or []:
        content = (window.get("note") or {}).get("content", "").strip()
        if not content:
            continue
        note_snapshot[window["code"]] = content
        any_note = True
        document.add_heading(window["label"], level=2)
        for text in content.splitlines():
            if text.strip():
                document.add_paragraph(text.strip())
    selected_note = (selected.get("note") or {}).get("content", "").strip()
    if selected_note and selected_note not in note_snapshot.values():
        note_snapshot[selected["code"]] = selected_note
        any_note = True
        document.add_heading(selected["label"], level=2)
        for text in selected_note.splitlines():
            if text.strip():
                document.add_paragraph(text.strip())
    if not any_note:
        document.add_paragraph("本周期尚未填写人工记录。")

    document.add_heading("六、风险与下一步", level=1)
    _add_analysis_section(document, "主要风险", analysis.get("risks"))
    _add_analysis_section(document, "下一步建议", analysis.get("actions"))
    _add_analysis_section(document, "需要的支持", analysis.get("support_needed"))

    document.add_heading("七、产品范围", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for index, value in enumerate(["产品", "ASIN", "MSKU", "店铺", "国家"]):
        table.rows[0].cells[index].text = value
    for product in module["products"]:
        cells = table.add_row().cells
        cells[0].text = str(product.get("product_name") or "-")
        cells[1].text = str(product.get("asin") or "-")
        cells[2].text = str(product.get("msku") or "-")
        cells[3].text = str(product.get("store_name") or "-")
        cells[4].text = str(product.get("country") or "-")

    document.save(path)
    archive_report(
        product_line=product_line,
        report_type=report_type,
        period_start=selected["start"],
        period_end=selected["end"],
        analysis_source=analysis_result.source,
        analysis_warning=analysis_result.warning,
        data_snapshot={"module": module, "selected": selected},
        target_snapshot=selected.get("targets") or [],
        note_snapshot=note_snapshot,
        analysis=analysis,
        artifact_path=path,
        db_path=database,
    )
    return path


def _add_summary_table(document: Document, window: dict[str, Any]) -> None:
    metrics = [
        ("销售额", "sales_amount", "currency"),
        ("销量", "units", "count"),
        ("利润", "profit", "currency"),
        ("利润率", "profit_margin", "ratio"),
        ("广告花费", "ad_spend", "currency"),
        ("广告销售额", "ad_sales", "currency"),
        ("TACOS", "tacos", "ratio"),
        ("CTR", "ctr", "ratio"),
        ("CPC", "cpc", "currency"),
        ("CVR", "cvr", "ratio"),
        ("退款金额", "refund_amount", "currency"),
        ("FBA可售库存", "fba_available", "count"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for index, value in enumerate(["指标", "本期", "对比上个等长周期"]):
        table.rows[0].cells[index].text = value
    summary = window.get("summary") or {}
    comparisons = window.get("comparisons") or {}
    for label, code, unit in metrics:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = _format_metric(summary.get(code), unit)
        cells[2].text = _format_change(comparisons.get(code))


def _add_target_table(document: Document, targets: list[dict[str, Any]]) -> None:
    if not targets:
        document.add_paragraph("当前报告周期尚未导入对应目标。")
        return
    document.add_heading("目标完成情况", level=2)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for index, value in enumerate(["指标", "实际", "目标", "完成/使用率", "状态", "目标范围"]):
        table.rows[0].cells[index].text = value
    for target in targets:
        cells = table.add_row().cells
        unit = str(target.get("unit") or "currency")
        cells[0].text = str(target.get("label") or target.get("metric_code"))
        cells[1].text = _format_metric(target.get("actual"), unit)
        cells[2].text = _format_metric(target.get("target"), unit)
        cells[3].text = _format_metric(target.get("completion"), "ratio")
        cells[4].text = _target_status_label(str(target.get("status") or "reference"))
        cells[5].text = str(target.get("scope_note") or "-")


def _add_bullets(document: Document, value: Any) -> None:
    items = value if isinstance(value, list) else ([value] if value else [])
    if not items:
        document.add_paragraph("暂无。")
        return
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def _add_analysis_section(document: Document, heading: str, value: Any) -> None:
    document.add_heading(heading, level=2)
    _add_bullets(document, value)


def _fallback_analysis(
    module: dict[str, Any], selected: dict[str, Any]
) -> dict[str, list[str]]:
    summary = selected.get("summary") or {}
    targets = selected.get("targets") or []
    notes = [
        (item.get("note") or {}).get("content", "").strip()
        for item in module.get("windows") or []
        if (item.get("note") or {}).get("content", "").strip()
    ]
    executive = [
        f"本期销售额：{_format_metric(summary.get('sales_amount'), 'currency')}；"
        f"销量：{_format_metric(summary.get('units'), 'count')}。"
    ]
    if summary.get("ad_spend") is not None:
        executive.append(
            f"广告花费：{_format_metric(summary.get('ad_spend'), 'currency')}；"
            f"TACOS：{_format_metric(summary.get('tacos'), 'ratio')}。"
        )
    if targets:
        target_lines = [
            f"{item['label']}目标完成/使用率为{_format_metric(item.get('completion'), 'ratio')}"
            for item in targets
            if item.get("completion") is not None
        ]
        executive.extend(target_lines[:3])
    if selected.get("warnings"):
        executive.append("存在数据缺失，相关结论需要结合后续同步结果复核。")
    return {
        "executive_summary": executive,
        "performance": ["当前为规则版分析，数字、目标完成率和趋势均由程序计算。"],
        "causes": notes[:3] or ["尚未填写人工原因判断。"],
        "risks": selected.get("warnings") or ["暂无程序识别出的数据风险。"],
        "actions": notes[-3:] or ["尚未填写下一步行动。"],
        "support_needed": ["尚未填写所需支持。"],
    }


def _format_metric(value: Any, unit: str) -> str:
    if value is None:
        return "数据缺失"
    number = float(value)
    if unit == "ratio":
        return f"{number * 100:.2f}%"
    if unit == "count":
        return f"{number:,.0f}"
    return f"{number:,.2f}"


def _format_change(value: Any) -> str:
    if value is None:
        return "无法比较"
    number = float(value)
    sign = "+" if number > 0 else ""
    return f"{sign}{number * 100:.2f}%"


def _target_status_label(value: str) -> str:
    return {
        "achieved": "已达到",
        "in_progress": "进行中",
        "exceeded": "超过上限",
        "within_budget": "预算内",
        "over_budget": "超预算",
        "missing": "实际数据缺失",
        "reference": "参考值",
    }.get(value, value)
