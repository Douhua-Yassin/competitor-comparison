"""Local Amazon US competitor monitor using an existing Chrome SellerSprite session."""
from __future__ import annotations

import asyncio
import json
import re
import sqlite3
import threading
import urllib.error
import urllib.request
from contextlib import asynccontextmanager, contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterator, Optional

from bs4 import BeautifulSoup
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "monitor.db"
INPUT_PATH = ROOT / "产品输入表.xlsx"
ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")
CDP_ENDPOINT = "http://127.0.0.1:9222"
SELLER_SPRITE_WAIT_SECONDS = 30
DELIVERY_ZIP_CODE = "90210"
SUCCESS_STATES = {"success", "partial_success"}
DB_TIMEOUT_SECONDS = 15
CDP_CONNECT_TIMEOUT_MS = 120_000

WRITE_LOCK = threading.RLock()
_workbook_signature: Optional[tuple[str, int, int]] = None
_last_sync_result: Optional[dict[str, Any]] = None

status: dict[str, Any] = {
    "running": False,
    "message": "idle",
    "success": 0,
    "partial": 0,
    "failed": 0,
    "total": 0,
    "current_asin": None,
    "browser_status": "unknown",
    "seller_sprite_status": "unknown",
    "delivery_status": "unknown",
    "last_sync": None,
    "sync_error": None,
}


class WorkbookSyncError(RuntimeError):
    pass


class DatabaseWriteError(RuntimeError):
    pass


def reset_sync_cache() -> None:
    global _workbook_signature, _last_sync_result
    _workbook_signature = None
    _last_sync_result = None


@contextmanager
def connection() -> Iterator[sqlite3.Connection]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH, timeout=DB_TIMEOUT_SECONDS)
    db.row_factory = sqlite3.Row
    db.execute(f"PRAGMA busy_timeout = {DB_TIMEOUT_SECONDS * 1000}")
    db.execute("PRAGMA foreign_keys = ON")
    try:
        with db:
            yield db
    finally:
        db.close()


def init_db() -> None:
    with WRITE_LOCK:
        with connection() as db:
            db.execute("PRAGMA journal_mode = WAL")
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS products (
                  asin TEXT PRIMARY KEY,
                  brand TEXT,
                  size_raw TEXT,
                  size_normalized TEXT,
                  is_self BOOLEAN,
                  input_rating_text TEXT,
                  input_rating_value REAL,
                  input_review_count INTEGER,
                  input_price REAL,
                  enabled BOOLEAN NOT NULL DEFAULT 1,
                  created_at DATETIME NOT NULL,
                  updated_at DATETIME NOT NULL
                );

                CREATE TABLE IF NOT EXISTS product_lines (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  name TEXT NOT NULL UNIQUE,
                  sheet_order INTEGER NOT NULL,
                  active BOOLEAN NOT NULL DEFAULT 1,
                  created_at DATETIME NOT NULL,
                  updated_at DATETIME NOT NULL
                );

                CREATE TABLE IF NOT EXISTS product_line_items (
                  product_line_id INTEGER NOT NULL,
                  asin TEXT NOT NULL,
                  brand TEXT,
                  size_raw TEXT,
                  size_normalized TEXT,
                  is_self BOOLEAN NOT NULL DEFAULT 0,
                  input_rating_text TEXT,
                  input_rating_value REAL,
                  input_review_count INTEGER,
                  input_price REAL,
                  item_order INTEGER NOT NULL,
                  active BOOLEAN NOT NULL DEFAULT 1,
                  created_at DATETIME NOT NULL,
                  updated_at DATETIME NOT NULL,
                  PRIMARY KEY(product_line_id, asin),
                  FOREIGN KEY(product_line_id) REFERENCES product_lines(id),
                  FOREIGN KEY(asin) REFERENCES products(asin)
                );

                CREATE TABLE IF NOT EXISTS crawl_records (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  asin TEXT NOT NULL,
                  record_date DATE NOT NULL,
                  captured_at DATETIME NOT NULL,
                  price REAL,
                  price_text TEXT,
                  sales_text TEXT,
                  rank_text TEXT,
                  rating_value REAL,
                  review_count INTEGER,
                  crawl_status TEXT NOT NULL,
                  error_message TEXT,
                  source TEXT NOT NULL CHECK(source IN ('live','fixture')),
                  data_source TEXT,
                  seller_sprite_status TEXT,
                  delivery_status TEXT,
                  FOREIGN KEY(asin) REFERENCES products(asin)
                );

                CREATE INDEX IF NOT EXISTS idx_crawl_asin_date
                  ON crawl_records(asin, record_date, captured_at);
                CREATE INDEX IF NOT EXISTS idx_line_items_active
                  ON product_line_items(product_line_id, active, item_order);
                """
            )
            existing = {row[1] for row in db.execute("PRAGMA table_info(crawl_records)")}
            for definition in (
                "data_source TEXT",
                "seller_sprite_status TEXT",
                "delivery_status TEXT",
            ):
                if definition.split()[0] not in existing:
                    db.execute("ALTER TABLE crawl_records ADD COLUMN " + definition)


def normalize_size(value: Any) -> Optional[str]:
    if value is None or not str(value).strip():
        return None
    text = str(value).strip()
    text = re.sub(r"(?i)(?<=\d)\s*(?:feet|foot|ft)\b", "", text)
    text = re.sub(r"[＊*×xX]", "×", text)
    parts = [part.strip() for part in text.split("×") if part.strip()]
    return "×".join(parts) + " ft" if parts else None


def parse_rating(value: Any) -> tuple[Optional[str], Optional[float], Optional[int]]:
    if value is None or not str(value).strip():
        return None, None, None
    text = str(value).strip()
    nums = re.findall(r"\d+(?:\.\d+)?", text.replace(",", ""))
    return (
        text,
        float(nums[0]) if nums else None,
        int(float(nums[1])) if len(nums) > 1 else None,
    )


def as_price(value: Any) -> Optional[float]:
    found = re.search(r"\d+(?:,\d{3})*(?:\.\d+)?", str(value or ""))
    return float(found.group().replace(",", "")) if found else None


def parse_is_self(value: Any) -> bool:
    return str(value or "").strip().lower() in {"是", "yes", "true", "1", "y"}


def workbook_signature(path: Path) -> tuple[str, int, int]:
    if not path.exists():
        raise WorkbookSyncError(f"未找到产品输入表：{path.name}")
    stat = path.stat()
    return (str(path.resolve()), stat.st_mtime_ns, stat.st_size)


def _read_workbook(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise WorkbookSyncError(f"未找到产品输入表：{path.name}")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookSyncError(f"无法读取产品输入表：{exc}") from exc

    lines: list[dict[str, Any]] = []
    try:
        for sheet_order, ws in enumerate(workbook.worksheets):
            items_by_asin: dict[str, dict[str, Any]] = {}
            for row_number, row in enumerate(
                ws.iter_rows(min_col=1, max_col=6, values_only=True), start=1
            ):
                asin = str(row[2] or "").strip().upper()
                if not ASIN_RE.fullmatch(asin):
                    continue
                rating_text, rating_value, review_count = parse_rating(row[3])
                existing = items_by_asin.get(asin)
                item_order = existing["item_order"] if existing else row_number
                items_by_asin[asin] = {
                    "asin": asin,
                    "brand": str(row[0] or "").strip() or None,
                    "size_raw": str(row[1] or "").strip() or None,
                    "size_normalized": normalize_size(row[1]),
                    "is_self": parse_is_self(row[5]),
                    "input_rating_text": rating_text,
                    "input_rating_value": rating_value,
                    "input_review_count": review_count,
                    "input_price": as_price(row[4]),
                    "item_order": item_order,
                }
            items = sorted(items_by_asin.values(), key=lambda item: item["item_order"])
            if items:
                lines.append(
                    {
                        "name": ws.title.strip() or f"Sheet{sheet_order + 1}",
                        "sheet_order": sheet_order,
                        "items": items,
                    }
                )
    finally:
        workbook.close()
    return lines


def _read_stable_workbook(path: Path) -> tuple[tuple[str, int, int], list[dict[str, Any]]]:
    first_signature = workbook_signature(path)
    lines = _read_workbook(path)
    second_signature = workbook_signature(path)
    if first_signature == second_signature:
        return second_signature, lines
    lines = _read_workbook(path)
    return workbook_signature(path), lines


def sync_workbook(path: Optional[Path] = None, force: bool = False) -> dict[str, Any]:
    global _workbook_signature, _last_sync_result

    workbook_path = path or INPUT_PATH
    current_signature = workbook_signature(workbook_path)
    if not force and current_signature == _workbook_signature and _last_sync_result:
        return dict(_last_sync_result)

    signature, lines = _read_stable_workbook(workbook_path)
    now = datetime.now().isoformat(timespec="seconds")
    unique_items: dict[str, dict[str, Any]] = {}
    for line in lines:
        for item in line["items"]:
            unique_items.setdefault(item["asin"], item)

    with WRITE_LOCK:
        init_db()
        try:
            with connection() as db:
                db.execute("UPDATE product_lines SET active=0, updated_at=?", (now,))
                db.execute("UPDATE product_line_items SET active=0, updated_at=?", (now,))
                db.execute("UPDATE products SET enabled=0, updated_at=?", (now,))

                for asin, item in unique_items.items():
                    db.execute(
                        """
                        INSERT INTO products (
                          asin, brand, size_raw, size_normalized, is_self,
                          input_rating_text, input_rating_value, input_review_count,
                          input_price, enabled, created_at, updated_at
                        ) VALUES (?,?,?,?,?,?,?,?,?,1,?,?)
                        ON CONFLICT(asin) DO UPDATE SET
                          brand=excluded.brand,
                          size_raw=excluded.size_raw,
                          size_normalized=excluded.size_normalized,
                          is_self=excluded.is_self,
                          input_rating_text=excluded.input_rating_text,
                          input_rating_value=excluded.input_rating_value,
                          input_review_count=excluded.input_review_count,
                          input_price=excluded.input_price,
                          enabled=1,
                          updated_at=excluded.updated_at
                        """,
                        (
                            asin,
                            item["brand"],
                            item["size_raw"],
                            item["size_normalized"],
                            item["is_self"],
                            item["input_rating_text"],
                            item["input_rating_value"],
                            item["input_review_count"],
                            item["input_price"],
                            now,
                            now,
                        ),
                    )

                for line in lines:
                    db.execute(
                        """
                        INSERT INTO product_lines (name, sheet_order, active, created_at, updated_at)
                        VALUES (?, ?, 1, ?, ?)
                        ON CONFLICT(name) DO UPDATE SET
                          sheet_order=excluded.sheet_order,
                          active=1,
                          updated_at=excluded.updated_at
                        """,
                        (line["name"], line["sheet_order"], now, now),
                    )
                    line_id = db.execute(
                        "SELECT id FROM product_lines WHERE name=?", (line["name"],)
                    ).fetchone()[0]
                    for item in line["items"]:
                        db.execute(
                            """
                            INSERT INTO product_line_items (
                              product_line_id, asin, brand, size_raw, size_normalized,
                              is_self, input_rating_text, input_rating_value,
                              input_review_count, input_price, item_order, active,
                              created_at, updated_at
                            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,1,?,?)
                            ON CONFLICT(product_line_id, asin) DO UPDATE SET
                              brand=excluded.brand,
                              size_raw=excluded.size_raw,
                              size_normalized=excluded.size_normalized,
                              is_self=excluded.is_self,
                              input_rating_text=excluded.input_rating_text,
                              input_rating_value=excluded.input_rating_value,
                              input_review_count=excluded.input_review_count,
                              input_price=excluded.input_price,
                              item_order=excluded.item_order,
                              active=1,
                              updated_at=excluded.updated_at
                            """,
                            (
                                line_id,
                                item["asin"],
                                item["brand"],
                                item["size_raw"],
                                item["size_normalized"],
                                item["is_self"],
                                item["input_rating_text"],
                                item["input_rating_value"],
                                item["input_review_count"],
                                item["input_price"],
                                item["item_order"],
                                now,
                                now,
                            ),
                        )
        except sqlite3.Error as exc:
            raise DatabaseWriteError(f"同步产品表时数据库写入失败：{exc}") from exc

    result = {
        "product_lines": len(lines),
        "items": sum(len(line["items"]) for line in lines),
        "unique_asins": len(unique_items),
        "synced_at": now,
    }
    _workbook_signature = signature
    _last_sync_result = dict(result)
    status.update(last_sync=now, sync_error=None)
    return result


def sync_or_raise(force: bool = False) -> dict[str, Any]:
    try:
        return sync_workbook(force=force)
    except WorkbookSyncError as exc:
        status["sync_error"] = str(exc)
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except DatabaseWriteError as exc:
        status["sync_error"] = str(exc)
        raise HTTPException(status_code=503, detail=str(exc)) from exc


def parse_amazon_html(html: str) -> dict[str, Any]:
    soup = BeautifulSoup(html, "html.parser")
    body = soup.get_text(" ", strip=True)
    lower = body.lower()
    if "captcha" in lower or "robot check" in lower:
        return {"crawl_status": "blocked_or_captcha", "error_message": "Amazon CAPTCHA"}
    if "page not found" in lower or "sorry! we couldn" in lower:
        return {"crawl_status": "page_not_found", "error_message": "Amazon page not found"}

    def first(selectors: list[str]) -> Optional[str]:
        for selector in selectors:
            node = soup.select_one(selector)
            if node and node.get_text(" ", strip=True):
                return node.get_text(" ", strip=True)
        return None

    price_text = first(
        [
            ".priceToPay .a-offscreen",
            "#corePrice_feature_div .a-offscreen",
            "#priceblock_ourprice",
            "#priceblock_dealprice",
            ".a-price .a-offscreen",
        ]
    )
    rating_text = first(["#acrPopover", "[data-hook='rating-out-of-text']"])
    reviews_text = first(["#acrCustomerReviewText", "[data-hook='total-review-count']"])
    detail = first(
        [
            "#productDetails_detailBullets_sections1",
            "#detailBullets_feature_div",
            "#productDetails_db_sections",
        ]
    )
    ranks = re.findall(r"#[\d,]+\s+in\s+[^#\n]{1,120}", detail or body, re.I)
    rating_num = re.search(r"(\d(?:\.\d)?)\s+out of 5", rating_text or body, re.I)
    review_num = re.search(r"[\d,]+", reviews_text or "")
    result = {
        "price": as_price(price_text),
        "price_text": price_text,
        "sales_text": first(
            [
                "#social-proofing-faceout-title-tk_bought",
                "#social-proofing-faceout-title",
                "#social-proofing-as-title",
            ]
        ),
        "rank_text": ranks[-1].strip() if ranks else None,
        "rating_value": float(rating_num.group(1)) if rating_num else None,
        "review_count": int(review_num.group().replace(",", "")) if review_num else None,
        "error_message": None,
    }
    available = any(
        result.get(key) is not None
        for key in ("price", "sales_text", "rank_text", "rating_value", "review_count")
    )
    result["crawl_status"] = (
        "success"
        if result["price"] is not None
        else ("partial_success" if available else "failed")
    )
    if not available:
        result["error_message"] = "Amazon public data was unavailable"
    return result


def parse_seller_sprite_html(html: str, asin: str) -> dict[str, Any]:
    soup = BeautifulSoup(html, "html.parser")
    root = soup.select_one(f'[name="seller-sprite-extension-quick-view-{asin}"]')
    if root is None:
        return {"seller_sprite_status": "unavailable"}

    def field(*labels: str) -> Optional[str]:
        for label in labels:
            normalized = label.rstrip(":：")
            for node in root.select(".word-title"):
                node_label = node.get_text(" ", strip=True).rstrip(":：")
                if node_label == normalized:
                    text = node.parent.get_text(" ", strip=True)
                    text = re.sub(
                        rf"^{re.escape(node.get_text(' ', strip=True))}\s*", "", text, count=1
                    ).strip(" :：")
                    return text or None
        return None

    root_text = root.get_text(" ", strip=True)
    rating = field("评分(评分数)")
    rating_value: Optional[float] = None
    review_count: Optional[int] = None
    if rating:
        match = re.search(r"(\d(?:\.\d+)?)\s*[（(]\s*([\d,]+)", rating)
        if match:
            rating_value = float(match.group(1))
            review_count = int(match.group(2).replace(",", ""))

    ranks: list[dict[str, Any]] = []
    for item in root.select("p.bsr-list-item"):
        text = item.get_text(" ", strip=True)
        match = re.search(r"#([\d,]+)\s+in\s+(.+)", text, re.I)
        if match:
            ranks.append(
                {
                    "rank": int(match.group(1).replace(",", "")),
                    "category": match.group(2).strip(),
                    "text": f"#{match.group(1)} in {match.group(2).strip()}",
                }
            )
    if not ranks:
        for match in re.finditer(
            r"#([\d,]+)\s+in\s+(.+?)(?=\s+#|\s+近30天|$)", root_text, re.I
        ):
            ranks.append(
                {
                    "rank": int(match.group(1).replace(",", "")),
                    "category": match.group(2).strip(),
                    "text": f"#{match.group(1)} in {match.group(2).strip()}",
                }
            )

    sales = field("近30天销量(父体)")
    return {
        "seller_sprite_status": "loaded",
        "asin": asin,
        "brand": field("品牌"),
        "price": as_price(field("价格")),
        "price_text": field("价格"),
        "rating_value": rating_value,
        "review_count": review_count,
        "parent_sales": None if not sales or sales.upper() == "N/A" else sales,
        "rank_list": ranks,
        "rank_text": ranks[-1]["text"] if ranks else None,
        "size": field("Size"),
    }


def join_messages(*messages: Optional[str]) -> Optional[str]:
    unique: list[str] = []
    for message in messages:
        if not message:
            continue
        for part in str(message).split("; "):
            if part and part not in unique:
                unique.append(part)
    return "; ".join(unique) if unique else None


def merge_sources(amazon: dict[str, Any], sprite: dict[str, Any]) -> dict[str, Any]:
    result = dict(amazon)
    sprite_loaded = sprite.get("seller_sprite_status") == "loaded"
    amazon_available = any(
        amazon.get(key) is not None
        for key in ("price", "sales_text", "rank_text", "rating_value", "review_count")
    )

    warnings: list[str] = []
    if sprite_loaded:
        mapping = {
            "price": "price",
            "price_text": "price_text",
            "rating_value": "rating_value",
            "review_count": "review_count",
            "parent_sales": "sales_text",
            "rank_text": "rank_text",
        }
        for sprite_key, result_key in mapping.items():
            value = sprite.get(sprite_key)
            if value is not None:
                old = result.get(result_key)
                if old is not None and old != value:
                    warnings.append(f"source_conflict:{result_key}")
                result[result_key] = value
        result["seller_sprite_status"] = "loaded"
        result["data_source"] = "seller_sprite+amazon" if amazon_available else "seller_sprite"
    else:
        result["seller_sprite_status"] = "unavailable"
        result["data_source"] = "amazon_public" if amazon_available else "none"
        warnings.append("seller_sprite_unavailable")

    available = any(
        result.get(key) is not None
        for key in ("price", "sales_text", "rank_text", "rating_value", "review_count")
    )
    if (
        result.get("price") is not None
        and result.get("sales_text") is not None
        and result.get("rank_text") is not None
    ):
        result["crawl_status"] = "success"
    elif available:
        result["crawl_status"] = "partial_success"
    else:
        result["crawl_status"] = amazon.get("crawl_status", "failed")

    result["error_message"] = join_messages(amazon.get("error_message"), *warnings)
    if result["crawl_status"] in SUCCESS_STATES and not result["error_message"]:
        result["error_message"] = None
    return result


async def set_delivery_location(page: Any) -> str:
    try:
        location = page.locator("#glow-ingress-line2")
        current = await location.inner_text(timeout=5000)
        if DELIVERY_ZIP_CODE in current or "china" not in current.lower():
            return "delivery_already_set"
        opener = page.locator("#nav-global-location-popover-link")
        if await opener.count() == 0:
            opener = page.locator("#glow-ingress-block")
        await opener.click(timeout=5000)
        zip_input = page.locator("#GLUXZipUpdateInput")
        await zip_input.wait_for(state="visible", timeout=10000)
        await zip_input.fill(DELIVERY_ZIP_CODE)
        await page.locator("#GLUXZipUpdate input.a-button-input").click(timeout=5000)
        await page.wait_for_function(
            """zip => {
              const success = document.querySelector('#GLUXHiddenSuccessSelectedAddressPlaceholder');
              const location = document.querySelector('#glow-ingress-line2');
              return (success && success.textContent.includes(zip)) ||
                (location && (!location.textContent.includes('China') || location.textContent.includes(zip)));
            }""",
            DELIVERY_ZIP_CODE,
            timeout=15000,
        )
        close = page.locator("#GLUXConfirmClose")
        if await close.count() and await close.is_visible():
            await close.click()
        await page.reload(wait_until="domcontentloaded", timeout=30000)
        return "delivery_set_90210"
    except Exception as exc:
        return "delivery_location_failed: " + str(exc)[:180]


def failed_result(code: str, exc: Exception | str) -> dict[str, Any]:
    return {
        "crawl_status": "failed",
        "error_message": f"{code}: {str(exc)[:500]}",
        "seller_sprite_status": "unavailable",
        "data_source": "none",
    }


async def crawl_product(asin: str, context: Any, delivery_state: dict[str, str]) -> dict[str, Any]:
    captured_at = datetime.now().isoformat(timespec="seconds")
    page = None
    result: dict[str, Any]
    try:
        try:
            page = await context.new_page()
        except Exception as exc:
            result = failed_result("page_create_failed", exc)
        else:
            try:
                await page.goto(
                    f"https://www.amazon.com/dp/{asin}",
                    wait_until="domcontentloaded",
                    timeout=60000,
                )
            except Exception as exc:
                result = failed_result("page_load_failed", exc)
            else:
                if delivery_state["value"] == "unknown":
                    delivery_state["value"] = await set_delivery_location(page)
                    status["delivery_status"] = delivery_state["value"]

                try:
                    amazon = parse_amazon_html(await page.content())
                except Exception as exc:
                    result = failed_result("parse_failed", exc)
                else:
                    sprite: dict[str, Any] = {"seller_sprite_status": "unavailable"}
                    selector = f'[name="seller-sprite-extension-quick-view-{asin}"]'
                    try:
                        root = page.locator(selector)
                        await root.wait_for(
                            state="attached", timeout=SELLER_SPRITE_WAIT_SECONDS * 1000
                        )
                        sprite = parse_seller_sprite_html(
                            await root.evaluate("el => el.outerHTML"), asin
                        )
                    except Exception:
                        sprite = {"seller_sprite_status": "unavailable"}

                    result = merge_sources(amazon, sprite)
                    status["seller_sprite_status"] = result.get(
                        "seller_sprite_status", "unknown"
                    )
    except Exception as exc:
        result = failed_result("crawl_failed", exc)
    finally:
        if page is not None:
            try:
                await page.close()
            except Exception:
                pass

    result.update(
        asin=asin,
        captured_at=captured_at,
        delivery_status=delivery_state["value"],
    )
    try:
        save_record(result, "live")
        result["persisted"] = True
    except DatabaseWriteError as exc:
        result["persisted"] = False
        result["crawl_status"] = "failed"
        result["error_message"] = join_messages(
            result.get("error_message"), f"database_write_failed: {exc}"
        )
    return result


def save_record(record: dict[str, Any], source: str) -> None:
    with WRITE_LOCK:
        try:
            with connection() as db:
                db.execute(
                    """
                    INSERT INTO crawl_records (
                      asin, record_date, captured_at, price, price_text, sales_text,
                      rank_text, rating_value, review_count, crawl_status,
                      error_message, source, data_source, seller_sprite_status,
                      delivery_status
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        record["asin"],
                        record["captured_at"][:10],
                        record["captured_at"],
                        record.get("price"),
                        record.get("price_text"),
                        record.get("sales_text"),
                        record.get("rank_text"),
                        record.get("rating_value"),
                        record.get("review_count"),
                        record["crawl_status"],
                        record.get("error_message"),
                        source,
                        record.get("data_source"),
                        record.get("seller_sprite_status"),
                        record.get("delivery_status"),
                    ),
                )
        except sqlite3.Error as exc:
            raise DatabaseWriteError(str(exc)) from exc


def get_cdp_info(endpoint: str = CDP_ENDPOINT) -> Optional[dict[str, Any]]:
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(endpoint + "/json/version", timeout=2) as response:
            if response.status != 200:
                return None
            payload = json.load(response)
    except (urllib.error.URLError, TimeoutError, OSError, ValueError):
        return None
    browser_name = str(payload.get("Browser", ""))
    if not re.search(r"(?:Chrome|Chromium)/", browser_name, re.I):
        return None
    if not payload.get("webSocketDebuggerUrl"):
        return None
    return payload


def check_cdp_available(endpoint: str = CDP_ENDPOINT) -> bool:
    return get_cdp_info(endpoint) is not None


def active_asins() -> list[str]:
    with connection() as db:
        return [
            row[0]
            for row in db.execute("SELECT asin FROM products WHERE enabled=1 ORDER BY asin")
        ]


async def connect_browser(playwright: Any) -> Any:
    info = get_cdp_info()
    if info is None:
        raise RuntimeError("9222 未返回有效的 Chrome 调试信息")
    websocket_url = str(info["webSocketDebuggerUrl"])
    return await playwright.chromium.connect_over_cdp(
        websocket_url, timeout=CDP_CONNECT_TIMEOUT_MS
    )


async def do_crawl(asins: list[str]) -> None:
    status.update(
        running=True,
        message="正在连接插件浏览器（首次连接最多等待 120 秒）",
        success=0,
        partial=0,
        failed=0,
        total=len(asins),
        current_asin=None,
        browser_status="connecting",
        seller_sprite_status="unknown",
        delivery_status="unknown",
    )
    delivery_state = {"value": "unknown"}
    try:
        from playwright.async_api import async_playwright

        async with async_playwright() as playwright:
            try:
                browser = await connect_browser(playwright)
                if not browser.contexts:
                    raise RuntimeError("插件浏览器没有可用的浏览器上下文")
            except Exception as exc:
                status.update(
                    browser_status="connection_failed",
                    message="插件浏览器连接失败：" + str(exc)[:300],
                )
                return

            context = browser.contexts[0]
            status["browser_status"] = "connected"
            for asin in asins:
                status.update(current_asin=asin, message=f"正在抓取 {asin}")
                try:
                    result = await crawl_product(asin, context, delivery_state)
                except Exception as exc:
                    status["failed"] += 1
                    status["message"] = f"{asin} 抓取失败：crawl_failed: {str(exc)[:180]}"
                    continue
                if result["crawl_status"] == "success":
                    status["success"] += 1
                elif result["crawl_status"] == "partial_success":
                    status["partial"] += 1
                else:
                    status["failed"] += 1
    except Exception as exc:
        if status["browser_status"] != "connected":
            status.update(
                browser_status="connection_failed",
                message="插件浏览器连接失败：" + str(exc)[:300],
            )
        else:
            status["message"] = "抓取任务异常：" + str(exc)[:300]
    finally:
        status.update(running=False, current_asin=None)
        if status["browser_status"] == "connected":
            status.update(
                browser_status="connected",
                message=(
                    f"抓取完成：成功 {status['success']}，"
                    f"部分成功 {status['partial']}，失败 {status['failed']}"
                ),
            )


def _line_rows(db: sqlite3.Connection) -> list[dict[str, Any]]:
    return [
        dict(row)
        for row in db.execute(
            "SELECT id, name, sheet_order FROM product_lines WHERE active=1 ORDER BY sheet_order, id"
        )
    ]


def table_data(range_name: str, product_line_id: Optional[int] = None) -> dict[str, Any]:
    with connection() as db:
        lines = _line_rows(db)
        if not lines:
            return {
                "product_lines": [],
                "selected_product_line_id": None,
                "dates": [],
                "products": [],
                "status": status,
            }
        valid_ids = {line["id"] for line in lines}
        selected_line_id = product_line_id if product_line_id in valid_ids else lines[0]["id"]
        products = [
            dict(row)
            for row in db.execute(
                """
                SELECT pli.asin, pli.brand, pli.size_raw, pli.size_normalized,
                       pli.is_self, pli.input_rating_text, pli.input_rating_value,
                       pli.input_review_count, pli.input_price, pli.item_order
                FROM product_line_items pli
                WHERE pli.product_line_id=? AND pli.active=1
                ORDER BY pli.item_order, pli.asin
                """,
                (selected_line_id,),
            )
        ]
        asins = [product["asin"] for product in products]
        if asins:
            placeholders = ",".join("?" for _ in asins)
            dates = [
                row[0]
                for row in db.execute(
                    f"""
                    SELECT DISTINCT record_date FROM crawl_records
                    WHERE asin IN ({placeholders})
                    ORDER BY record_date DESC
                    """,
                    asins,
                )
            ]
            records = [
                dict(row)
                for row in db.execute(
                    f"""
                    SELECT * FROM crawl_records
                    WHERE asin IN ({placeholders})
                    ORDER BY captured_at DESC, id DESC
                    """,
                    asins,
                )
            ]
        else:
            dates = []
            records = []

    if range_name in {"3d", "7d"}:
        minimum = (date.today() - timedelta(days=int(range_name[0]) - 1)).isoformat()
        dates = [record_date for record_date in dates if record_date >= minimum]

    cells: dict[tuple[str, str], dict[str, Any]] = {}
    for product in products:
        for record_date in dates:
            rows = [
                row
                for row in records
                if row["asin"] == product["asin"] and row["record_date"] == record_date
            ]
            good = [row for row in rows if row["crawl_status"] in SUCCESS_STATES]
            if good:
                cells[(product["asin"], record_date)] = good[0]
            elif rows:
                cells[(product["asin"], record_date)] = rows[0]

    baseline: dict[tuple[Optional[str], str], float] = {}
    for product in products:
        if not product["is_self"]:
            continue
        for record_date in dates:
            record = cells.get((product["asin"], record_date))
            if record is None or record["price"] is None:
                continue
            key = (product["size_normalized"], record_date)
            baseline[key] = min(baseline.get(key, record["price"]), record["price"])

    for product in products:
        product["cells"] = []
        for record_date in dates:
            record = cells.get((product["asin"], record_date))
            tone = "neutral"
            base = baseline.get((product["size_normalized"], record_date))
            if (
                record
                and not product["is_self"]
                and record["price"] is not None
                and base is not None
            ):
                if record["price"] < base:
                    tone = "low"
                elif record["price"] > base:
                    tone = "high"
            product["cells"].append({"record": record, "tone": tone})

    return {
        "product_lines": lines,
        "selected_product_line_id": selected_line_id,
        "dates": dates,
        "products": products,
        "status": status,
    }


@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()
    try:
        sync_workbook(force=True)
    except (WorkbookSyncError, DatabaseWriteError) as exc:
        status["sync_error"] = str(exc)
        status["message"] = str(exc)
    yield


app = FastAPI(title="Amazon Competitor Monitor", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=ROOT / "app" / "static"), name="static")
templates = Jinja2Templates(directory=ROOT / "app" / "templates")


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    try:
        sync_workbook()
    except (WorkbookSyncError, DatabaseWriteError) as exc:
        status["sync_error"] = str(exc)
    return templates.TemplateResponse(request, "index.html", {})


@app.post("/api/sync")
def api_sync():
    return sync_or_raise(force=True)


@app.post("/api/crawl")
async def api_crawl():
    if status["running"]:
        raise HTTPException(status_code=409, detail="抓取任务正在运行")
    sync_or_raise()
    asins = active_asins()
    if not asins:
        raise HTTPException(status_code=400, detail="产品输入表中没有有效 ASIN")
    if not check_cdp_available():
        status.update(
            browser_status="plugin_browser_not_started",
            message="请先启动插件浏览器",
            sync_error=None,
        )
        raise HTTPException(status_code=503, detail="请先启动插件浏览器")

    # Reserve the batch before yielding back to the event loop. Without this,
    # two near-simultaneous requests can both schedule a crawl before do_crawl
    # has a chance to set status["running"].
    status.update(
        running=True,
        message="抓取任务已排队",
        total=len(asins),
        current_asin=None,
    )
    try:
        asyncio.create_task(do_crawl(asins))
    except Exception:
        status.update(running=False, message="抓取任务启动失败")
        raise
    return {"started": True, "total": len(asins)}


@app.get("/api/status")
def api_status():
    return status


@app.get("/api/table")
def api_table(range: str = "all", product_line_id: Optional[int] = None):
    if range not in {"all", "3d", "7d"}:
        raise HTTPException(status_code=400, detail="range must be all, 3d, or 7d")
    sync_or_raise()
    return table_data(range, product_line_id)
